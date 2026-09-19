"""Build a Google-Docs-friendly .docx (RTL, header/footer, no overlapping frames)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "REPORT_FA.md"
DOCX = ROOT / "REPORT_FA.docx"
LOGO = ROOT / "assets" / "iut-logo.png"
NAVY = RGBColor(0x14, 0x3A, 0x6B)
GOLD = RGBColor(0xB8, 0x8A, 0x2E)


def bidi_p(p):
    pPr = p._element.get_or_add_pPr()
    el = OxmlElement("w:bidi")
    pPr.append(el)
    return p


def rtl_run(run):
    rPr = run._element.get_or_add_rPr()
    el = OxmlElement("w:rtl")
    rPr.append(el)
    rPr.append(OxmlElement("w:cs"))
    return run


def set_run_font(run, name="Tahoma", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    rtl_run(run)


def add_page_number(paragraph):
    run1 = paragraph.add_run("صفحه ")
    set_run_font(run1, size=9, color=NAVY)
    fld1 = paragraph.add_run()
    set_run_font(fld1, size=9, color=NAVY)
    r = fld1._element
    fc1 = OxmlElement("w:fldChar")
    fc1.set(qn("w:fldCharType"), "begin")
    r.append(fc1)
    run2 = paragraph.add_run()
    set_run_font(run2, size=9, color=NAVY)
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = " PAGE "
    run2._element.append(t)
    run3 = paragraph.add_run()
    set_run_font(run3, size=9, color=NAVY)
    fc2 = OxmlElement("w:fldChar")
    fc2.set(qn("w:fldCharType"), "end")
    run3._element.append(fc2)
    run4 = paragraph.add_run(" از ")
    set_run_font(run4, size=9, color=NAVY)
    run5 = paragraph.add_run()
    set_run_font(run5, size=9, color=NAVY)
    r5 = run5._element
    fc3 = OxmlElement("w:fldChar")
    fc3.set(qn("w:fldCharType"), "begin")
    r5.append(fc3)
    run6 = paragraph.add_run()
    set_run_font(run6, size=9, color=NAVY)
    t2 = OxmlElement("w:instrText")
    t2.set(qn("xml:space"), "preserve")
    t2.text = " NUMPAGES "
    run6._element.append(t2)
    run7 = paragraph.add_run()
    set_run_font(run7, size=9, color=NAVY)
    fc4 = OxmlElement("w:fldChar")
    fc4.set(qn("w:fldCharType"), "end")
    run7._element.append(fc4)


def setup_section(section, first_page=False):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    h = section.header.paragraphs[0]
    bidi_p(h)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = h.add_run("دانشگاه صنعتی اصفهان  ·  پروژه کارشناسی  ·  سید محمد فاطمی")
    set_run_font(r, size=9, color=NAVY, bold=True)
    pPr = h._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "C4A35A")
    pBdr.append(bottom)
    pPr.append(pBdr)

    f = section.footer.paragraphs[0]
    bidi_p(f)
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = f.add_run("استاد راهنما: دکتر زینب مالکی    |    ")
    set_run_font(r, size=9, color=NAVY)
    add_page_number(f)
    r = f.add_run("    |    ۱۴۰۵")
    set_run_font(r, size=9, color=NAVY)
    pPr = f._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "8")
    top.set(qn("w:space"), "4")
    top.set(qn("w:color"), "C4A35A")
    pBdr.append(top)
    pPr.append(pBdr)

    # empty first-page header/footer so cover is clean
    fh = section.first_page_header.paragraphs[0]
    fh.text = ""
    ff = section.first_page_footer.paragraphs[0]
    ff.text = ""


def add_runs_inline(p, text, size=11, bold=False, color=None):
    # split **bold**, `code`
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color or NAVY)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=size - 1, color=color)
        else:
            clean = part.replace("*", "")
            run = p.add_run(clean)
            set_run_font(run, size=size, bold=bold, color=color)


def cover(doc: Document):
    if LOGO.exists():
        p = doc.add_paragraph()
        bidi_p(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Cm(4.2))

    def center(text, size, bold=False, space_after=6, color=NAVY):
        p = doc.add_paragraph()
        bidi_p(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
        return p

    center("بسمه تعالی", 14, True, 10)
    center("دانشگاه صنعتی اصفهان", 18, True, 4)
    center("دانشکده مهندسی برق و کامپیوتر", 13, False, 2)
    center("گروه مهندسی کامپیوتر", 13, False, 14)
    center("پروژهٔ کارشناسی", 14, True, 16)
    center("طبقه‌بندی مبتنی بر گراف با", 16, True, 2)
    center("استفاده از لم نظم زمردی", 16, True, 2)
    center("(Szemerdi)", 16, True, 18)
    center("دانشجو: سید محمد فاطمی", 13, False, 4)
    center("استاد راهنما: دکتر زینب مالکی", 13, False, 4)
    center("دانشیار، دانشکده مهندسی برق و کامپیوتر", 11, False, 16)
    center("اصفهان — ۱۴۰۵ هجری شمسی", 12, True, 0)
    doc.add_page_break()


def toc_placeholder(doc: Document):
    p = doc.add_paragraph()
    bidi_p(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("فهرست مطالب")
    set_run_font(r, size=16, bold=True, color=NAVY)
    n = doc.add_paragraph()
    bidi_p(n)
    add_runs_inline(
        n,
        "پس از باز کردن این فایل در Google Docs: Insert → Table of contents "
        "(با لینک). تیترها از سبک Heading ساخته شده‌اند.",
        size=10,
    )
    doc.add_page_break()


def apply_heading_style(p, level):
    p.style = f"Heading {level}"
    bidi_p(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color=NAVY)


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            bidi_p(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            val = row[j] if j < len(row) else ""
            val = val.replace("**", "")
            r = p.add_run(val)
            set_run_font(r, size=8, bold=(i == 0), color=None if i else RGBColor(0xFF, 0xFF, 0xFF))
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "143A6B")
                shd.set(qn("w:val"), "clear")
                cell._tePr = cell._tc.get_or_add_tcPr()
                cell._tc.get_or_add_tcPr().append(shd)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.add_paragraph()


def parse_md_tables_and_blocks(md: str):
    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    blocks = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+:?", lines[i + 1].replace("|", " | ")):
            table = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^:?-+:?$", raw[0].replace(" ", "")):
                    table.append(raw)
                i += 1
            blocks.append(("table", table))
            continue
        blocks.append(("line", line))
        i += 1
    return blocks


def body_from_md(doc: Document, md: str):
    for kind, payload in parse_md_tables_and_blocks(md):
        if kind == "table":
            add_table(doc, payload)
            continue
        line = payload
        s = line.strip()
        if not s:
            continue
        if s == "---":
            continue
        if s.startswith("# "):
            p = doc.add_paragraph(s[2:])
            apply_heading_style(p, 1)
            continue
        if s.startswith("## "):
            p = doc.add_paragraph(s[3:])
            apply_heading_style(p, 2)
            continue
        if s.startswith("### "):
            p = doc.add_paragraph(s[4:])
            apply_heading_style(p, 3)
            continue
        if s.startswith("> "):
            p = doc.add_paragraph()
            bidi_p(p)
            p.paragraph_format.right_indent = Cm(0.4)
            add_runs_inline(p, s[2:], size=11)
            continue
        if s.startswith("- ") or s.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            bidi_p(p)
            add_runs_inline(p, s[2:], size=11)
            continue
        if re.match(r"^\d+\.\s", s):
            p = doc.add_paragraph(style="List Number")
            bidi_p(p)
            add_runs_inline(p, re.sub(r"^\d+\.\s", "", s), size=11)
            continue
        p = doc.add_paragraph()
        bidi_p(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        add_runs_inline(p, s, size=11)


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Tahoma"
    style.font.size = Pt(11)
    setup_section(doc.sections[0])
    cover(doc)
    toc_placeholder(doc)
    md = MD.read_text(encoding="utf-8")
    body_from_md(doc, md)
    doc.save(DOCX)
    print(f"wrote {DOCX} ({DOCX.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
