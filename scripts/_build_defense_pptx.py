"""BSc defense slides (widescreen PPTX, RTL Persian)."""
from __future__ import annotations

import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "THESIS_FA_slides.pptx"
OUT_ALT = ROOT / "THESIS_FA_slides_new.pptx"
LOGO = ROOT / "assets" / "iut-logo.png"
FIG = ROOT / "assets" / "figures"

PAPER = RGBColor(0xEB, 0xE4, 0xD6)
PAPER2 = RGBColor(0xF7, 0xF1, 0xE6)
PANEL = RGBColor(0xFF, 0xFC, 0xF8)
INK = RGBColor(0x1B, 0x18, 0x16)
MUTED = RGBColor(0x6B, 0x64, 0x5C)
LINE = RGBColor(0xD4, 0xCB, 0xBA)
ACCENT = RGBColor(0xA3, 0x3B, 0x24)
NAVY = RGBColor(0x1A, 0x33, 0x4A)
NAVY_DEEP = RGBColor(0x0D, 0x1C, 0x2C)
SAGE = RGBColor(0x2C, 0x5E, 0x52)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SHELL = RGBColor(0xDE, 0xD5, 0xC4)
FONT_FA = "B Nazanin"
FONT_EN = "Calibri"
FA_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")
ARABIC = re.compile(r"[\u0600-\u06FF]")
LATIN = re.compile(
    r"[A-Za-z\u00C0-\u024F][A-Za-z\u00C0-\u024F0-9+\-_.\u00B2\u00B3\u00B9]*"
)


def _clean_fa(text: str) -> str:
    return (
        text.replace("\u2066", "")
        .replace("\u2069", "")
        .replace("\u200e", "")
        .replace("\u200f", "")
        .replace("ۀ", "ه")
        .replace("ٔ", "")
        .replace("·", " | ")
    )


def _style_run(run, *, rtl: bool) -> None:
    font = FONT_FA if rtl else FONT_EN
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    rPr.set("rtl", "1" if rtl else "0")
    rPr.set("lang", "fa-IR" if rtl else "en-US")
    pairs = (("ea", FONT_FA if rtl else FONT_EN), ("cs", FONT_FA if rtl else FONT_EN), ("latin", FONT_EN))
    for tag, face in pairs:
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", face)


def _p(p, align="r") -> None:
    pPr = p._p.get_or_add_pPr()
    if align == "r":
        p.alignment = PP_ALIGN.RIGHT
        pPr.set("rtl", "1")
        pPr.set("algn", "r")
    elif align == "en-r":
        p.alignment = PP_ALIGN.RIGHT
        pPr.set("rtl", "0")
        pPr.set("algn", "r")
    elif align == "c":
        p.alignment = PP_ALIGN.CENTER
        pPr.set("rtl", "0")
        pPr.set("algn", "ctr")
    else:
        p.alignment = PP_ALIGN.LEFT
        pPr.set("rtl", "0")
        pPr.set("algn", "l")


def _wipe(p) -> None:
    el = p._p
    for child in list(el):
        if child.tag == qn("a:r"):
            el.remove(child)


def add_run(p, text, size=20, bold=False, color=INK, rtl=True) -> None:
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    _style_run(run, rtl=rtl)


def add_mixed(p, text, size=20, bold=False, color=INK) -> None:
    text = _clean_fa(text)
    if not ARABIC.search(text):
        add_run(p, text, size=size, bold=bold, color=color, rtl=False)
        return
    pos = 0
    for m in LATIN.finditer(text):
        left = text[m.start() - 1] if m.start() else " "
        right = text[m.end()] if m.end() < len(text) else " "
        glued = left in "-/" or right in "-/" or bool(ARABIC.match(left) or ARABIC.match(right))
        if glued:
            continue
        if m.start() > pos:
            add_run(p, text[pos:m.start()], size=size, bold=bold, color=color, rtl=True)
        add_run(p, m.group(), size=size, bold=bold, color=color, rtl=False)
        pos = m.end()
    if pos < len(text):
        add_run(p, text[pos:], size=size, bold=bold, color=color, rtl=True)


def box(slide, l, t, w, h, fill=None, line=None, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, l, t, w, h)
    if rounded:
        try:
            sh.adjustments[0] = 0.1
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    return sh


def _spPr(shape):
    return shape._element.spPr


def fill_alpha(shape, pct: float) -> None:
    solid = _spPr(shape).find(qn("a:solidFill"))
    if solid is None:
        return
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        return
    el = srgb.find(qn("a:alpha"))
    if el is None:
        el = etree.SubElement(srgb, qn("a:alpha"))
    el.set("val", str(int(pct * 1000)))


def shadow(
    shape,
    *,
    blur=0.32,
    dist=0.09,
    dir_deg=120,
    alpha=22,
    color="1A334A",
) -> None:
    spPr = _spPr(shape)
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    lst = etree.SubElement(spPr, qn("a:effectLst"))
    shdw = etree.SubElement(lst, qn("a:outerShdw"))
    shdw.set("blurRad", str(int(blur * 914400)))
    shdw.set("dist", str(int(dist * 914400)))
    shdw.set("dir", str(int(dir_deg * 60000)))
    shdw.set("algn", "ctr")
    shdw.set("rotWithShape", "0")
    srgb = etree.SubElement(shdw, qn("a:srgbClr"))
    srgb.set("val", color)
    a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int(alpha * 1000)))


def gradient(shape, c1: RGBColor, c2: RGBColor, angle: float = 108) -> None:
    shape.fill.gradient()
    shape.fill.gradient_angle = angle
    stops = shape.fill.gradient_stops
    stops[0].color.rgb = c1
    stops[1].color.rgb = c2


def blob(slide, l, t, w, h, color, alpha=10):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, w, h)
    sh.line.fill.background()
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    fill_alpha(sh, alpha)
    return sh


def paper_bg(prs, slide, *, inverse=False):
    W, H = prs.slide_width, prs.slide_height
    bg = box(slide, 0, 0, W, H, PAPER)
    if inverse:
        gradient(bg, NAVY_DEEP, NAVY, 118)
    else:
        gradient(bg, PAPER, PAPER2, 108)
        blob(slide, Inches(9.8), Inches(-1.4), Inches(5.2), Inches(5.2), ACCENT, 8)
        blob(slide, Inches(-1.6), Inches(4.6), Inches(4.4), Inches(4.4), NAVY, 9)
    rail = box(slide, 0, 0, Inches(0.14), H, ACCENT if not inverse else ACCENT)
    return bg


def tf(slide, l, t, w, h):
    frame = slide.shapes.add_textbox(l, t, w, h).text_frame
    frame.word_wrap = True
    return frame


def write(frame, lines, size=20, bold=False, color=INK, space=8, align="r", bullet=False):
    first = True
    for line in lines:
        p = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        _wipe(p)
        if ARABIC.search(line):
            use = align
        elif align == "c":
            use = "c"
        elif align == "r":
            use = "en-r"
        else:
            use = "l"
        _p(p, use)
        p.space_after = Pt(space)
        if bullet:
            pPr = p._p.get_or_add_pPr()
            bu_font = etree.SubElement(pPr, qn("a:buFont"))
            bu_font.set("typeface", "Calibri")
            bu_char = etree.SubElement(pPr, qn("a:buChar"))
            bu_char.set("char", "•")
        add_mixed(p, line, size=size, bold=bold, color=color)


def hairline(slide, l, t, w):
    box(slide, l, t, w, Pt(1.1), LINE)


def new(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def chrome(prs, slide, title: str, page: int, kicker: str = ""):
    W, H = prs.slide_width, prs.slide_height
    paper_bg(prs, slide)
    bar = box(slide, Inches(0.45), Inches(0.28), Inches(12.45), Inches(0.92), PANEL, rounded=True)
    shadow(bar, blur=0.22, dist=0.05, alpha=16)
    f = tf(slide, Inches(0.7), Inches(0.42), W - Inches(1.6), Inches(0.62))
    write(f, [title], size=26, bold=True, color=NAVY, space=0)
    f = tf(slide, Inches(0.55), H - Inches(0.38), Inches(8.5), Inches(0.28))
    write(f, ["خوشه‌بندی گرافی با لم نظم زمردی"], size=11, color=MUTED, space=0)
    chip = box(slide, W - Inches(1.55), H - Inches(0.42), Inches(0.95), Inches(0.28), NAVY, rounded=True)
    f = tf(slide, W - Inches(1.55), H - Inches(0.44), Inches(0.95), Inches(0.28))
    write(f, [str(page).translate(FA_DIGITS)], size=11, color=WHITE, space=0, align="c")


def add_picture_fit(slide, path: Path, l, t, w, h):
    if not path.exists():
        f = tf(slide, l, t, w, Inches(0.5))
        write(f, [f"شکل پیدا نشد: {path.name}"], size=16, color=MUTED, space=0)
        return
    pic = slide.shapes.add_picture(str(path), l, t)
    nat_w, nat_h = int(pic.width), int(pic.height)
    scale = min(w / nat_w, h / nat_h)
    pic.width = int(nat_w * scale)
    pic.height = int(nat_h * scale)
    pic.left = int(l + (w - pic.width) / 2)
    pic.top = int(t + (h - pic.height) / 2)


def fig_slide(prs, slides, title, path, caption="", kicker="شکل", big=False):
    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, title, len(slides), kicker=kicker)
    lines = list(caption) if isinstance(caption, (list, tuple)) else ([caption] if caption else [])
    pic_w = 9.15 if big else 8.05
    shell = box(sl, Inches(0.36), Inches(1.28), Inches(pic_w), Inches(5.7), SHELL, rounded=True)
    shadow(shell, blur=0.22, dist=0.06, alpha=16)
    box(sl, Inches(0.46), Inches(1.38), Inches(pic_w - 0.2), Inches(5.5), PANEL, rounded=True)
    add_picture_fit(
        sl,
        path,
        Inches(0.58),
        Inches(1.48),
        Inches(pic_w - 0.44),
        Inches(5.3),
    )
    if lines:
        text_x = 0.36 + pic_w + 0.16
        text_w = 13.0 - text_x
        card = box(sl, Inches(text_x), Inches(1.28), Inches(text_w), Inches(5.7), PANEL, rounded=True)
        shadow(card, blur=0.22, dist=0.06, alpha=16)
        box(sl, Inches(12.74), Inches(1.28), Inches(0.1), Inches(5.7), ACCENT)
        f = tf(sl, Inches(text_x + 0.18), Inches(1.48), Inches(text_w - 0.42), Inches(5.3))
        f._txBody.find(qn("a:bodyPr")).set("anchor", "ctr")
        write(f, lines, size=18, color=INK, space=12, bullet=True)
    return sl


def trio(slide, cards, top=1.7, height=4.6):
    n = len(cards)
    gap = 0.22
    width = (12.3 - gap * (n - 1)) / n
    for i, (kick, title, body) in enumerate(cards):
        x = Inches(0.5 + (n - 1 - i) * (width + gap))
        card = box(slide, x, Inches(top), Inches(width), Inches(height), PANEL, rounded=True)
        shadow(card, blur=0.28, dist=0.08, alpha=18)
        box(slide, x, Inches(top), Inches(width), Inches(0.12), ACCENT if i == 0 else NAVY)
        f = tf(slide, x + Inches(0.25), Inches(top + 0.35), Inches(width - 0.5), Inches(0.4))
        write(f, [kick], size=14, color=ACCENT, space=0)
        f = tf(slide, x + Inches(0.25), Inches(top + 0.9), Inches(width - 0.5), Inches(1.3))
        write(f, [title], size=22 if n > 3 else 26, bold=True, color=NAVY, space=0)
        f = tf(slide, x + Inches(0.25), Inches(top + 2.3), Inches(width - 0.5), Inches(2.0))
        write(f, [body], size=18, color=INK, space=0)


def numbered(slide, items, top=1.38):
    n = max(len(items), 1)
    row_h = min(1.02, 5.55 / n)
    for i, item in enumerate(items):
        y = Inches(top) + Inches(i * row_h)
        h = Inches(row_h - 0.1)
        card = box(slide, Inches(0.5), y, Inches(12.35), h, PANEL, rounded=True)
        shadow(card, blur=0.2, dist=0.05, alpha=14)
        box(slide, Inches(12.68), y + Inches(0.12), Inches(0.1), h - Inches(0.24), ACCENT if i == 0 else NAVY)
        f = tf(slide, Inches(11.55), y + Inches(0.14), Inches(1.0), h - Inches(0.22))
        write(f, [f"{i+1:02d}".translate(FA_DIGITS)], size=16, bold=True, color=ACCENT, space=0)
        f = tf(slide, Inches(0.75), y + Inches(0.14), Inches(10.65), h - Inches(0.22))
        write(f, [item], size=18, color=INK, space=0)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    slides = []

    s = new(prs)
    slides.append(s)
    paper_bg(prs, s)
    plate = box(s, Inches(0.5), Inches(0.42), Inches(12.35), Inches(4.7), PANEL, rounded=True)
    shadow(plate, blur=0.45, dist=0.14, alpha=28)
    if LOGO.exists():
        s.shapes.add_picture(str(LOGO), Inches(11.55), Inches(0.62), height=Inches(0.95))
    f = tf(s, Inches(0.85), Inches(0.7), Inches(10.2), Inches(0.4))
    write(f, ["دانشگاه صنعتی اصفهان  |  دانشکده مهندسی برق و کامپیوتر"], size=15, color=ACCENT, space=0)
    f = tf(s, Inches(0.85), Inches(1.45), Inches(11.6), Inches(2.15))
    write(
        f,
        ["طبقه‌بندی مبتنی بر گراف", "با استفاده از لم نظم زمردی"],
        size=36,
        bold=True,
        color=NAVY,
        space=6,
    )
    f = tf(s, Inches(0.85), Inches(3.85), Inches(10), Inches(0.4))
    write(f, ["Szemerédi regularity lemma"], size=16, color=MUTED, space=0)
    box(s, Inches(0.85), Inches(4.45), Inches(2.4), Inches(0.08), ACCENT)
    meta = [
        ("دانشجو", "سید محمد فاطمی"),
        ("سال", "۱۴۰۵"),
        ("مقطع", "کارشناسی مهندسی کامپیوتر"),
        ("استاد راهنما", "دکتر زینب مالکی"),
    ]
    for i, (lab, val) in enumerate(meta):
        x = Inches(0.5 + (len(meta) - 1 - i) * 3.2)
        card = box(s, x, Inches(5.4), Inches(3.02), Inches(1.55), PANEL, rounded=True)
        shadow(card, blur=0.22, dist=0.06, alpha=18)
        f = tf(s, x + Inches(0.18), Inches(5.55), Inches(2.7), Inches(0.32))
        write(f, [lab], size=12, color=ACCENT, space=0)
        f = tf(s, x + Inches(0.18), Inches(5.9), Inches(2.7), Inches(0.75))
        write(f, [val], size=16, bold=True, color=INK, space=0)

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "اهمیت و بیان مسئله", len(slides), kicker="مسئله")
    lines = [
        "خیلی از داده‌ها برچسب آماده ندارند. خوشه‌بندی برای همین مهم است: نمونه‌های شبیه باید خودشان یک گروه شوند، مثل مشتری‌های هم‌سلیقه یا تصویرهای هم‌خانواده.",
        "یک راه جاافتاده برای دیدن این شباهت، ساختن گراف است. هر نمونه یک رأس می‌شود و یال میان دو رأس می‌گوید چقدر به هم نزدیک‌اند.",
        "گرهٔ کار از همین‌جا شروع می‌شود. همین که تعداد نمونه‌ها بالا برود، گرافی که همسایگی را نشان می‌داد دیگر در حافظه و زمان جا نمی‌شود.",
    ]
    card = box(sl, Inches(6.35), Inches(1.36), Inches(6.45), Inches(5.55), PANEL, rounded=True)
    shadow(card, blur=0.28, dist=0.08, alpha=18)
    box(sl, Inches(12.7), Inches(1.36), Inches(0.1), Inches(5.55), ACCENT)
    f = tf(sl, Inches(6.6), Inches(1.55), Inches(5.9), Inches(5.15))
    write(f, lines, size=20, color=INK, space=16)
    shell = box(sl, Inches(0.48), Inches(1.36), Inches(5.65), Inches(5.55), SHELL, rounded=True)
    shadow(shell, blur=0.28, dist=0.08, alpha=18)
    box(sl, Inches(0.62), Inches(1.5), Inches(5.37), Inches(5.27), PANEL, rounded=True)
    add_picture_fit(sl, FIG / "concept_why_graph.png", Inches(0.95), Inches(1.85), Inches(4.7), Inches(4.55))

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "چالش‌ها", len(slides), kicker="مسئله")
    lines = [
        "برای هر دو نمونه یک عدد شباهت حساب می‌شود و همهٔ این عددها باید بمانند.",
        "با n نمونه، تعداد این عددها از مرتبهٔ n² است؛ حافظه خیلی زود پر می‌شود.",
        "بخش بزرگی از الگوریتم‌های خوشه‌بندی گرافی زمانی نزدیک n³ دارند.",
        "مثلاً برای حدود یازده هزار نمونه، ماتریس شباهت نزدیک صد و بیست میلیون خانه دارد.",
    ]
    card = box(sl, Inches(6.35), Inches(1.36), Inches(6.45), Inches(5.55), PANEL, rounded=True)
    shadow(card, blur=0.28, dist=0.08, alpha=18)
    box(sl, Inches(12.7), Inches(1.36), Inches(0.1), Inches(5.55), ACCENT)
    f = tf(sl, Inches(6.6), Inches(1.7), Inches(5.9), Inches(4.9))
    write(f, lines, size=20, color=INK, space=14, bullet=True)
    shell = box(sl, Inches(0.48), Inches(1.36), Inches(5.65), Inches(5.55), SHELL, rounded=True)
    shadow(shell, blur=0.28, dist=0.08, alpha=18)
    box(sl, Inches(0.62), Inches(1.5), Inches(5.37), Inches(5.27), PANEL, rounded=True)
    add_picture_fit(sl, FIG / "concept_similarity_matrix.png", Inches(0.95), Inches(1.85), Inches(4.7), Inches(4.55))

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "متودولوژی", len(slides), kicker="لم")
    lines = [
        "گراف شباهت یک بار ساخته می‌شود. هر دو راه آزمایش از همین گراف شروع می‌کنند.",
        "راه پایه: همان الگوریتم، مستقیم روی گراف اصلی.",
        "راه لم: تقسیم رأس‌ها به گروه‌های هم‌اندازه، خلاصهٔ رابطهٔ گروه‌ها با یک چگالی، بعد خوشه‌بندی روی گراف کوچک.",
        "برچسب گراف کوچک به نمونه‌ها برمی‌گردد. کیفیت هر دو راه با NMI و برچسب واقعی سنجیده می‌شود.",
    ]
    card = box(sl, Inches(6.35), Inches(1.36), Inches(6.45), Inches(5.55), PANEL, rounded=True)
    shadow(card, blur=0.28, dist=0.08, alpha=18)
    box(sl, Inches(12.7), Inches(1.36), Inches(0.1), Inches(5.55), ACCENT)
    f = tf(sl, Inches(6.6), Inches(1.7), Inches(5.9), Inches(4.9))
    write(f, lines, size=18, color=INK, space=12, bullet=True)
    shell = box(sl, Inches(0.48), Inches(1.36), Inches(5.65), Inches(5.55), SHELL, rounded=True)
    shadow(shell, blur=0.28, dist=0.08, alpha=18)
    box(sl, Inches(0.62), Inches(1.5), Inches(5.37), Inches(5.27), PANEL, rounded=True)
    add_picture_fit(sl, FIG / "concept_why_lemma.png", Inches(0.95), Inches(1.85), Inches(4.7), Inches(4.55))

    fig_slide(
        prs,
        slides,
        "قدم اول: گروه‌های هم‌اندازه",
        FIG / "concept_partition.png",
        [
            "لم اول رأس‌ها را به چند گروه تقریباً هم‌اندازه تقسیم می‌کند.",
            "لکه‌های رنگی همین گروه‌ها هستند.",
            "نقطه‌های خاکستری در هیچ گروه تمیزی جا نشده‌اند.",
            "آن ته مانده را مجموعهٔ استثنایی می‌نامند.",
        ],
        kicker="لم",
    )

    fig_slide(
        prs,
        slides,
        "جفت منظم: یال‌ها یکنواخت‌اند",
        FIG / "concept_regular.png",
        [
            "منظم بودن دربارهٔ یال بین دو گروه است، نه داخل هر گروه.",
            "باید زیرمجموعه‌های به‌قدر کافی بزرگ را هم چک کرد، نه تکه‌های خیلی کوچک را.",
            "حد اندازه، کسری از خود گروه است. از آن حد به بعد، چگالی هر دو زیرمجموعه باید نزدیک چگالی کل بماند.",
            "اگر این طور باشد، یک عدد چگالی کل رابطه را خلاصه می‌کند.",
        ],
        kicker="لم",
    )

    fig_slide(
        prs,
        slides,
        "جفت نامنظم: یک عدد کافی نیست",
        FIG / "concept_irregular.png",
        [
            "اینجا بالای تصویر پر از یال است و پایین تقریباً خالی.",
            "میانگین، این تفاوت را پنهان می‌کند.",
            "لم چنین جفتی را نامنظم می‌گیرد و باید بعداً بریده شود.",
        ],
        kicker="لم",
    )

    fig_slide(
        prs,
        slides,
        "اپسیلون، مرز منظم و نامنظم",
        FIG / "fig_eps_regular_pair.png",
        [
            "چپ جفت منظم است: چگالی تقریباً ثابت و نزدیک d.",
            "راست جفت نامنظم است: یک تکه شلوغ و بقیه خلوت.",
            "اپسیلون هم حد اختلاف چگالی است و هم کف اندازهٔ زیرمجموعه.",
            "زیرمجموعهٔ کوچک‌تر از این کف کنار گذاشته می‌شود. بزرگ‌تر از آن، باید چگالی‌اش نزدیک چگالی کل بماند.",
        ],
        kicker="لم",
        big=True,
    )

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "چطور جفت نامنظم را پیدا می‌کنیم", len(slides), kicker="افراز")
    cards = [
        ("آزمون ۱", "یال خیلی کم، جفت منظم است", "اگر میانگین یال بین دو گروه از حدی کمتر باشد، جفت را منظم می‌گیریم. گراف خالی هم منظم است.", FIG / "concept_alon1.png"),
        ("آزمون ۳", "همسایهٔ مشترکِ زیاد", "اگر دسته‌ای از رأس‌ها همسایه‌های خیلی شبیه داشته باشند، همان دسته شاهد نامنظمی است.", FIG / "concept_alon3.png"),
        ("آزمون ۲", "درجهٔ خیلی دور از میانگین", "رأس‌هایی که درجه‌شان از میانگین خیلی پرت است جمع می‌شوند. اگر تعدادشان زیاد باشد، certificate همان‌هاست.", FIG / "concept_alon2.png"),
    ]
    gap = 0.22
    width = (12.3 - gap * 2) / 3
    for i, (kick, title, body, path) in enumerate(cards):
        x = 0.5 + (2 - i) * (width + gap)
        card = box(sl, Inches(x), Inches(1.5), Inches(width), Inches(5.15), PANEL, rounded=True)
        shadow(card, blur=0.28, dist=0.08, alpha=18)
        box(sl, Inches(x), Inches(1.5), Inches(width), Inches(0.1), ACCENT if i == 0 else NAVY)
        f = tf(sl, Inches(x + 0.18), Inches(1.7), Inches(width - 0.36), Inches(0.32))
        write(f, [kick], size=14, color=ACCENT, space=0)
        f = tf(sl, Inches(x + 0.18), Inches(2.02), Inches(width - 0.36), Inches(0.7))
        write(f, [title], size=18, bold=True, color=NAVY, space=0)
        f = tf(sl, Inches(x + 0.18), Inches(2.7), Inches(width - 0.36), Inches(1.25))
        write(f, [body], size=14, color=INK, space=0)
        add_picture_fit(sl, path, Inches(x + 0.32), Inches(4.05), Inches(width - 0.64), Inches(2.25))

    fig_slide(
        prs,
        slides,
        "برش گروه، داخل ساخت افراز",
        FIG / "concept_refine.png",
        [
            "جای این کار در راه لم، بعد از آزمون جفت‌ها و قبل از ساخت گراف کوچک است.",
            "هنوز خوشه‌بندی نیست. گروه نامنظم دو تکه می‌شود تا بشود یال‌ها را با یک چگالی عوض کرد.",
            "تکهٔ قرمز همان certificate است و از گروه جدا می‌شود. بقیهٔ گروه سر جایش می‌ماند.",
            "خرده‌هایی که به هیچ تکه‌ای نمی‌خورند به مجموعهٔ استثنایی می‌روند.",
        ],
        kicker="افراز",
    )

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "سه اصلاح مقاله روی همین افراز", len(slides), kicker="افراز")
    numbered(
        sl,
        [
            "هر گروه حداکثر با یک گروه دیگر نامنظم حساب می‌شود. بدون این قید، تعداد گروه‌ها در هر دور چند برابر می‌شود.",
            "تکهٔ نامنظم با روش درجهٔ Fiorucci پیدا می‌شود، نه با جست‌وجوی کامل نسخهٔ اصلی Alon.",
            "وقتی گروه‌ها به‌قدر کافی کوچک شدند، حلقه می‌ایستد. نتیجه تقریباً منظم است، نه یک افراز اثبات‌شده.",
        ],
    )

    fig_slide(
        prs,
        slides,
        "از یال‌های بین دو گروه تا یک چگالی",
        FIG / "concept_density.png",
        [
            "چپ، همهٔ یال‌های بین دو گروه را نشان می‌دهد.",
            "راست، همان رابطه فقط با یک میله مانده است.",
            "ضخامت میله یعنی چگالی: سهم یال‌های موجود از همهٔ یال‌های ممکن.",
        ],
        kicker="لم",
    )

    fig_slide(
        prs,
        slides,
        "گراف کوچک: هر گروه یک رأس",
        FIG / "concept_reduced.png",
        [
            "هر گروه با فلش به یک رأس هم‌رنگ تبدیل می‌شود.",
            "آن رأس نمایندهٔ همهٔ نمونه‌های گروه است.",
            "ضخامت یال بین دو رأس، همان چگالی دو گروه است.",
            "خوشه‌بندی از اینجا به بعد روی همین گراف کوچک انجام می‌شود.",
        ],
        kicker="لم",
    )

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "همان گراف، یک بار مستقیم و یک بار با لم", len(slides), kicker="pipeline")
    paths = [
        ("راه پایه", "خوشه‌بندی مستقیم روی گراف اصلی. رنگ هر رأس همان خوشهٔ اوست.", FIG / "concept_path_direct.png"),
        ("راه لم", "رأس‌های زیاد به ۴ رأس کوچک می‌شوند. خوشه‌بندی روی همین گراف چهار رأسی انجام می‌شود.", FIG / "concept_path_lemma.png"),
    ]
    for i, (title, body, path) in enumerate(paths):
        x = 0.45 + (1 - i) * 6.45
        card = box(sl, Inches(x), Inches(1.48), Inches(6.2), Inches(5.2), PANEL, rounded=True)
        shadow(card, blur=0.28, dist=0.08, alpha=18)
        box(sl, Inches(x), Inches(1.48), Inches(6.2), Inches(0.1), ACCENT if i == 0 else NAVY)
        f = tf(sl, Inches(x + 0.25), Inches(1.68), Inches(5.7), Inches(0.45))
        write(f, [title], size=22, bold=True, color=NAVY, space=0)
        add_picture_fit(sl, path, Inches(x + 0.35), Inches(2.15), Inches(5.5), Inches(3.55))
        f = tf(sl, Inches(x + 0.25), Inches(5.75), Inches(5.7), Inches(0.8))
        write(f, [body], size=16, color=INK, space=0)

    fig_slide(
        prs,
        slides,
        "برگرداندن خوشه به تک‌تک نمونه‌ها",
        FIG / "concept_lift.png",
        [
            "خوشه‌بندی فقط به رأس‌های گراف کوچک رنگ داده است.",
            "هر نمونه، رنگ رأس گروهی را می‌گیرد که عضو آن است.",
            "پس خروجی باز هم یک برچسب برای هر نمونهٔ اصلی است.",
        ],
        kicker="pipeline",
    )

    fig_slide(
        prs,
        slides,
        "کل مسیر آزمایش، از داده تا NMI",
        FIG / "fig_pipeline.png",
        [
            "داده خوانده می‌شود، در صورت نیاز کوچک می‌شود، و گراف شباهت یک بار ساخته می‌شود.",
            "شاخهٔ پایه همان جا روی گراف اصلی خوشه می‌زند.",
            "شاخهٔ لم از افراز و گراف کوچک رد می‌شود و برچسب را به نمونه‌ها برمی‌گرداند.",
            "هر دو شاخه با برچسب واقعی و NMI مقایسه می‌شوند.",
        ],
        kicker="pipeline",
    )

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "چهار الگوریتم، روی هر دو راه", len(slides), kicker="pipeline")
    trio(
        sl,
        [
            ("۱", "SPC", "خوشه‌بندی طیفی. گراف را با بردارهای ویژه می‌بُرد."),
            ("۲", "APC", "انتشار وابستگی. نمونه‌ها خودشان نمایندهٔ خوشه را انتخاب می‌کنند."),
            ("۳", "DSet", "مجموعهٔ غالب. یک تکهٔ خیلی متراکم را به‌عنوان خوشه برمی‌دارد."),
            ("۴", "SPRG", "خوشه‌بندی طیفی با شباهت جنگل تصادفی. شباهت از جنگل می‌آید، بعد برش طیفی."),
        ],
        top=1.55,
        height=4.7,
    )

    fig_slide(
        prs,
        slides,
        "اصلاح آخر برچسب، بعد از برگشت به نمونه‌ها",
        FIG / "concept_polish.png",
        [
            "دایره گروه لم است، نه مرز خوشه.",
            "اول همهٔ اعضای یک گروه یک رنگ می‌گیرند، حتی اگر یکی‌شان میان همسایه‌های رنگ دیگر باشد.",
            "این مرحله فقط رنگ خوشهٔ همان نقطه را عوض می‌کند.",
            "عضویت او در گروه لم دست نمی‌خورد. این مرحله جزو خود لم نیست.",
        ],
        kicker="بعد از لم",
    )

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "تنظیم آزمایش و معیار مقایسه", len(slides), kicker="آزمایش")
    numbered(
        sl,
        [
            "بیست مجموعهٔ داده و چهار الگوریتم: هشتاد خانه، که هفتاد و هشت تایش معتبر ماند.",
            "ده مجموعه با همهٔ نمونه‌ها اجرا شد و ده مجموعه با حدود چهارصد نمونهٔ طبقه‌ای.",
            "برای هر خانه، شبکهٔ تنظیمات مثل آزمایش مقاله انتخاب شد و بهترین مقدار گزارش شد.",
            "ادعای اصلی این ارائه، مقایسه با راه پایهٔ خودمان است، نه رونویسی جدول چاپ‌شدهٔ مقاله.",
        ],
    )

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "۶۰ بهتر، ۳ برابر، ۱۵ بدتر", len(slides), kicker="آزمایش")
    f = tf(sl, Inches(0.55), Inches(1.45), Inches(12.2), Inches(0.7))
    write(f, ["هفتاد و هشت مقایسه. هر بار همان الگوریتم، یک بار روی گراف اصلی و یک بار روی گراف کوچک."], size=18, color=MUTED, space=0)
    cards = [
        ("۶۰", "گراف کوچک بهتر", SAGE),
        ("۳", "برابر", NAVY),
        ("۱۵", "گراف اصلی بهتر", ACCENT),
    ]
    for i, (n, lab, col) in enumerate(cards):
        x = Inches(0.5 + (2 - i) * 4.2)
        card = box(sl, x, Inches(2.3), Inches(3.95), Inches(3.15), PANEL, rounded=True)
        shadow(card, blur=0.32, dist=0.1, alpha=20)
        box(sl, x, Inches(2.3), Inches(3.95), Inches(0.12), col)
        f = tf(sl, x, Inches(2.6), Inches(3.95), Inches(1.5))
        write(f, [n], size=60, bold=True, color=col, space=0, align="c")
        f = tf(sl, x, Inches(4.3), Inches(3.95), Inches(0.7))
        write(f, [lab], size=20, color=INK, space=0, align="c")

    fig_slide(
        prs,
        slides,
        "نمودار اختلاف: بالای صفر یعنی لم بهتر است",
        FIG / "fig_delta_nmi_full.png",
        [
            "هر میله اختلاف NMI راه لم با راه پایه است.",
            "بالای صفر: لم بهتر بوده است.",
            "پایین صفر: راه پایه بهتر بوده است.",
            "NMI نزدیکی خوشهٔ پیش‌بینی‌شده به برچسب واقعی را می‌سنجد.",
        ],
        kicker="نتیجه",
    )

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "سه مجموعه با بیشترین بهبود", len(slides), kicker="مثال")
    trio(
        sl,
        [
            ("شراب", "پایه ضعیف، لم قوی", "سه الگوریتم از حدود نیم به نزدیک ۰/۸۹ رسیدند."),
            ("آپاندیسیت", "تطابق کامل", "دو الگوریتم از پایهٔ متوسط به ۱ رسیدند."),
            ("اسکناس", "طیفی ضعیف بود", "خوشه‌بندی طیفی از حدود ۰/۱۹ به حدود ۰/۸۲ رسید."),
        ],
        height=4.5,
    )

    sl = new(prs)
    slides.append(sl)
    chrome(prs, sl, "جمع‌بندی", len(slides), kicker="پایان")
    trio(
        sl,
        [
            ("۱", "لم فشرده می‌کند", "گراف بزرگ به چند ده رأس کاهش پیدا می‌کند."),
            ("۲", "کیفیت معمولاً بهتر ماند", "در ۶۰ مقایسه گراف کوچک بهتر بود، در ۳ مقایسه برابر، در ۱۵ مقایسه بدتر."),
            ("۳", "الگوریتم قدیمی قوی‌تر شد", "همان روش‌های کلاسیک، بعد از کوچک کردن گراف، در بیشتر آزمایش‌ها دقیق‌تر شدند. بیشترین سود جایی بود که راه پایه ناپایدار بود."),
        ],
        height=4.5,
    )

    s = new(prs)
    slides.append(s)
    paper_bg(prs, s, inverse=True)
    blob(s, Inches(9.2), Inches(-1.2), Inches(5.6), Inches(5.6), ACCENT, 18)
    blob(s, Inches(-1.8), Inches(4.2), Inches(4.8), Inches(4.8), SAGE, 14)
    plate = box(s, Inches(1.35), Inches(2.05), Inches(10.65), Inches(3.35), PANEL, rounded=True)
    shadow(plate, blur=0.5, dist=0.16, alpha=32, color="000000")
    f = tf(s, Inches(1.7), Inches(2.4), Inches(10), Inches(1.3))
    write(f, ["پرسش و پاسخ"], size=44, bold=True, color=NAVY, space=0)
    box(s, Inches(1.7), Inches(3.85), Inches(2.2), Inches(0.09), ACCENT)
    f = tf(s, Inches(1.7), Inches(4.15), Inches(10), Inches(0.6))
    write(f, ["با سپاس از دکتر زینب مالکی"], size=20, color=MUTED, space=0)

    try:
        prs.save(OUT)
        return OUT
    except PermissionError:
        prs.save(OUT_ALT)
        return OUT_ALT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path} ({path.stat().st_size} bytes)")
